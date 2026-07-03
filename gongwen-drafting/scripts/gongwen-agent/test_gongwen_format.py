#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import tempfile
import unittest
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))
from docx import Document
import gongwen_format as gf


class GongwenFormatTests(unittest.TestCase):
    def test_red_header_uses_literal_placeholders_when_missing(self):
        data = gf.HechaRequestData(
            issuing_org="测试公司",
            recipient_org="测试文物局",
            project_name="测试项目",
            date_text="2026年7月2日",
        )
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "hecha.docx"
            gf.create_hecha_docx(data, output)
            text = gf.collect_text(output)

        self.assertIn("【文号】", text)
        self.assertIn("签发人：【签发人】", text)

    def test_audit_fails_old_blank_red_header_placeholders(self):
        with tempfile.TemporaryDirectory() as tmp:
            output = Path(tmp) / "old-style.docx"
            doc = Document()
            gf.apply_document_geometry(doc)
            gf.add_red_header(doc, "测试公司", doc_number="文号", signer="")
            gf.add_title(doc, "测试公司关于办理测试项目用地范围内文物保护许可的请示")
            gf.add_body_paragraph(doc, "测试文物局：", first_line=False)
            gf.add_body_paragraph(doc, "项目面积1公顷。")
            gf.add_attachment_item(doc, 1, "测试附件")
            gf.add_signature(doc, "测试公司", "2026年7月2日")
            doc.save(output)

            errors = gf.audit_hecha_docx(output)

        self.assertIn("红头文号缺少“【文号】”占位或正式文号", errors)


if __name__ == "__main__":
    unittest.main()
